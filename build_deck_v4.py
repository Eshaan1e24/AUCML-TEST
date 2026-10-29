"""Build IHDS_poverty_escape_presentation_v4.pptx: a lean main talk (slides 1 to 32) and a clearly marked backup section (33 to 40).

Every number comes from the notebook outputs (outputs/key_numbers.json, the CSV files and chart_dump.pkl).
Run from the project folder (after the notebook and make_charts_v3.py):  python build_deck_v4.py
"""
import json
import pickle

import numpy as np
import pandas as pd
from pptx.dml.color import RGBColor

import deck_helpers as dh
from deck_helpers import (BLANK, DATA_SRC, H, INK, MSO_ANCHOR, MSO_SHAPE, MUTED, ORANGE, PRESENTER, SECTIONS, TEAL, TINT, WHITE,
                          Inches, Pt, box, count, new_slide, note, picture, prs, rect, stat, table, text)

NUM = dh.ROOT / "outputs"
k = json.load(open(NUM / "key_numbers.json"))
dump = pickle.load(open(NUM / "chart_dump.pkl", "rb"))
cv_summary = pd.read_csv(NUM / "cv_summary.csv").set_index("Model")
dictionary = pd.read_csv(NUM / "feature_dictionary.csv")
res = {r["Model"]: r for r in k["results"]}
fm = {r["Model"]: r for r in k["full_metrics"]}
ENS = "Soft-vote ensemble (LR+RF+XGB)"
HEAD = k["best_model"]
rep = {r["Model"]: r for r in k["repeated_splits"]}
pair = {(p["model_a"], p["model_b"]): p for p in k["paired_differences"]}
cal = {r["Model"]: r for r in k["calibration"]}
edu = k["education"]
EA, EB = edu["A: baseline characteristics"], edu["B: A + baseline poverty gap"]
pr = k["poverty_rates"]
fsel = pd.DataFrame(k["feature_selection"]["table"]).set_index("Features kept")
rv = k["random_vs_village"]
over = {r["Model"]: r for r in k["overfitting_check"]}
BACKUP_FILL = RGBColor(0xFF, 0xF1, 0xE6)
import subprocess, sys
_check = subprocess.run([sys.executable, str(dh.ROOT / 'quick_check.py')], capture_output=True, text=True, cwd=str(dh.ROOT))
QUICK_CHECK_LINES = ['$ python quick_check.py'] + [ln for ln in _check.stdout.splitlines() if ln.strip()][:12]
assert _check.returncode == 0 and 'OK' in _check.stdout, _check.stderr
chart_data = dump["chart_data"]
y_test, probs, thr = dump["y_test"], dump["test_probs"], dump["thresholds"]
lift = k["lift"]


def pct(x, d=0):
    return f"{x * 100:.{d}f}%"


def notes(slide, content):
    slide.notes_slide.notes_text_frame.text = content


def mono(slide, x, y, w, h, lines, size=12):
    rect(slide, x, y, w, h)
    tb = slide.shapes.add_textbox(x + Inches(0.12), y + Inches(0.08), w - Inches(0.24), h - Inches(0.16))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = line
        r.font.name, r.font.size, r.font.color.rgb = "Consolas", Pt(size), INK


def backup_slide(title):
    s = new_slide(None, title, DATA_SRC)
    tag = rect(s, Inches(10.9), Inches(0.32), Inches(2.1), Inches(0.42), fill=BACKUP_FILL, line=None, rnd=0.5)
    tf = tag.text_frame
    tf.paragraphs[0].alignment = dh.PP_ALIGN.CENTER
    run = tf.paragraphs[0].add_run()
    run.text = "Backup (if asked)"
    run.font.size, run.font.bold, run.font.name, run.font.color.rgb = Pt(12), True, dh.FONT, ORANGE
    return s


# =========================================================================================================== 1 Title
count[0] += 1
s = prs.slides.add_slide(BLANK)
s.background.fill.solid()
s.background.fill.fore_color.rgb = dh.BG
rect(s, 0, 0, Inches(4.2), H, fill=TEAL, line=None, shape=MSO_SHAPE.RECTANGLE)
text(s, Inches(0.5), Inches(0.6), Inches(3.4), Inches(1), ["Applications and Use Cases of Machine Learning (TE7483)", "ESE Case Study Project"], 14, WHITE, True)
text(s, Inches(0.5), Inches(4.8), Inches(3.4), Inches(2), ["SDG 1: No Poverty", "SDG 10: Reduced Inequalities"], 16, WHITE, True)
text(s, Inches(4.8), Inches(1.0), Inches(8.1), Inches(2.2), "Predicting Household Escape from Poverty in India", 38, INK, True)
text(s, Inches(4.8), Inches(2.55), Inches(8.1), Inches(1), "A Machine Learning Analysis of Longitudinal Socio-Economic Factors (India Human Development Survey, 2004-05 and 2011-12)", 18, MUTED)
rect(s, Inches(4.8), Inches(3.75), Inches(8.0), Inches(2.3))
text(s, Inches(5.05), Inches(3.9), Inches(7.6), Inches(2.1),
     ["**Team members", "Anhad Mahajan (PRN 24070126031)", "Aman Sumeet Saxena (PRN 24070126030)", "Eshaan Sumesh (PRN 24070126058)", "Gokul Sethumadhavan (PRN 24070126061)"], 16, gap=0.25)
text(s, Inches(4.8), Inches(6.2), Inches(8), Inches(1.0),
     ["Guide: Dr. Shrikrishna Kolhar", "AIML-A, B.Tech 2024-28, academic year 2026-27", "Symbiosis Institute of Technology, Pune"], 14, MUTED, gap=0.1)
notes(s, "Anhad: Hello everyone. We are Anhad, Aman, Eshaan and Gokul, and our project is about poverty in India. We asked a practical question: which poor households will still be poor seven years later? "
         "To answer it we used a real survey that visited the same households twice. We built and tested four machine learning models, and we made a small demo app.")

# =========================================================================================================== 2 Contents
s = new_slide(None, "Contents")
for i, sec in enumerate(SECTIONS):
    col, row = i // 4, i % 4
    x = Inches(0.6) + Inches(6.2) * col
    y = Inches(1.65) + Inches(1.28) * row
    rect(s, x, y, Inches(5.95), Inches(1.1))
    text(s, x + Inches(0.2), y, Inches(0.8), Inches(1.1), str(i + 1), 30, TEAL, True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + Inches(1.0), y + Inches(0.15), Inches(4.8), Inches(0.5), sec, 19, INK, True)
    text(s, x + Inches(1.0), y + Inches(0.58), Inches(4.8), Inches(0.4), "Presented by " + PRESENTER[i + 1], 13, MUTED)
notes(s, "Anhad: Our talk has eight parts. I will cover the problem, the earlier studies and the conclusion. Aman will explain the data and how we built and ran the project. "
         "Eshaan will explain the methods and what the model has learned. Gokul will present the results. All four of us can answer questions on every part.")

# =========================================================================================================== 3 Project on a page
s = new_slide(None, "The project on one page")
cards = [("Problem", "Which poor households will still be poor seven years later? If a programme knew this, it could help those households first."),
         ("Data", f"A survey visited the same households in 2004-05 and again in 2011-12. We study the {k['n_kept']:,} poor households that could be found again."),
         ("Method", f"We use {k['n_features']} facts about each household from 2004-05 to train four models. We test them on villages they have never seen."),
         ("Result", f"About {pct(k['escape_rate'])} of these households escaped poverty. Our best model, XGBoost, scores a ROC-AUC of {res[HEAD]['ROC_AUC']:.2f}, where 0.50 would be guessing."),
         ("What it means", f"The 20% of households that the model rates riskiest stayed poor {pct(lift[-1]['stayed_poor_rate'])} of the time. The average was {pct(1 - float(np.mean(y_test)))}, so the model is useful for ranking."),
         ("Limits", "The accuracy is modest and the data are old. The model shows patterns, and it should be used to screen groups of households, not to judge one family.")]
for i, (head, body) in enumerate(cards):
    x = Inches(0.6) + Inches(4.1) * (i % 3)
    y = Inches(1.6) + Inches(2.7) * (i // 3)
    box(s, x, y, Inches(3.95), Inches(2.5), head, [body], 15, ORANGE if head == "Limits" else TEAL)
notes(s, "If you remember only one slide, remember this one. It holds the question, the data, the method, the result, what the result means and what the model cannot do. "
         "Every other slide fills in one of these six boxes.")

# =========================================================================================================== 4 Problem and objectives (Anhad)
s = new_slide(1, "Problem statement and objectives")
box(s, Inches(0.6), Inches(1.6), Inches(5.9), Inches(2.4), "Problem",
    ["Welfare budgets are limited. We want to tell in advance which poor households are least likely to leave poverty on their own, so that help can reach them first."], 16)
box(s, Inches(0.6), Inches(4.2), Inches(5.9), Inches(2.6), "Question",
    ["Using only what is known when a household is poor, can we predict whether it is still poor seven years later? And which household conditions go together with escaping?"], 16)
box(s, Inches(6.8), Inches(1.6), Inches(6.0), Inches(5.2), "Objectives",
    ["1. Predict escape from poverty using 2004-05 information only.",
     "2. Compare four models on villages never used for training.",
     "3. Find the factors the models rely on, using SHAP.",
     "4. Check whether education goes together with escape.",
     "5. Build a small demo app."], 16, gap=0.6)
notes(s, "Anhad: Most poverty numbers describe a single year. A welfare programme, though, needs to know who is likely to stay poor. So we turned this into a prediction problem. "
         "We take a household that was poor in 2004-05 and ask whether it had escaped by 2011-12. Along the way we also wanted to learn which conditions matter, and to build a small demo that shows the idea.")

# =========================================================================================================== 5 Domain, users, relevance (Anhad)
s = new_slide(1, "Application domain, users and real-world relevance")
box(s, Inches(0.6), Inches(1.6), Inches(4.0), Inches(2.6), "Application domain",
    ["Our domain is social welfare and poverty targeting, which is part of public policy.", "The machine learning task is classification: did the household escape or stay poor?"], 14)
box(s, Inches(4.8), Inches(1.6), Inches(4.0), Inches(2.6), "Intended users",
    ["Government welfare departments can rank poor households by risk.", "NGOs and microfinance institutions can choose villages and families.", "Researchers can study what goes together with leaving poverty."], 14)
box(s, Inches(9.0), Inches(1.6), Inches(3.8), Inches(2.6), "Why it matters",
    ["Budgets are limited, so programmes must choose whom to help first.", "Machine learning is already used for this. In Togo it reduced the number of poor people left out by 4 to 21% compared with targeting by geography (Aiken et al., 2022)."], 14)
box(s, Inches(0.6), Inches(4.4), Inches(12.2), Inches(2.4), "Sustainable development goals",
    ["Our work supports SDG 1 (No Poverty), whose target 1.2 is to reduce poverty according to national definitions. We measure the money side of poverty, through household spending.",
     "It also supports SDG 10 (Reduced Inequalities), whose target 10.2 is to include all people, whatever their sex, origin, religion or other status. We use caste, religion and state as inputs so that differences between groups can be seen."], 15)
notes(s, "Anhad: Our domain is social welfare and poverty targeting. The people who could use this are welfare departments, NGOs and researchers. It matters because budgets are limited, "
         "and programmes need a way to decide whom to reach first. Machine learning is already used for this: a 2022 paper in Nature used it to target cash aid in Togo. Our work supports SDG 1 and SDG 10.")

# =========================================================================================================== 6 Registered idea to project (Anhad)
s = new_slide(1, "From the registered idea to this project")
box(s, Inches(0.6), Inches(1.6), Inches(5.9), Inches(3.3), "What we registered",
    ["Take households that were poor in the first survey round and predict whether they had crossed the poverty line by the second round.",
     "Compare Logistic Regression, Random Forest and XGBoost.",
     "Rank the factors with SHAP, and optionally add a small propensity-score analysis."], 15, gap=0.5)
box(s, Inches(6.8), Inches(1.6), Inches(6.0), Inches(3.3), "What we delivered and improved",
    ["We did all of this and added an ensemble model.",
     "We used one poverty line, the Tendulkar line, in both rounds.",
     "We tested on whole villages that the models had never seen.",
     "We checked that fewer features work almost as well, and we built a working demo app."], 15, gap=0.5)
note(s, Inches(0.6), Inches(5.2), Inches(12.2), Inches(1.5),
     "The project covers poor households in 2004-05 that we could follow up to 2011-12. It predicts risk and shows patterns. It does not prove what causes a household to escape.", 16)
notes(s, "Anhad: We registered a clear idea, and we have delivered all of it. We improved it in four ways: we used one poverty line in both rounds, we test on whole villages, "
         "we checked how many features are really needed, and we built the demo. The optional propensity-score idea became our education analysis. We say clearly that it shows a link and not a cause.")

# =========================================================================================================== 7 Literature (Anhad)
s = new_slide(2, "What earlier studies found")
table(s, Inches(0.6), Inches(1.5),
      [["Study", "Main finding", "How it connects to our project"],
       ["Krishna (2004)", "A study of 35 Rajasthan villages: 11.1% of households escaped poverty and 7.9% fell into it.", "It shows why escape is worth studying by itself."],
       ["Thorat et al. (2017)", "Same survey as ours: about 65% of poor households in 2005 had escaped by 2012. Adivasi households escaped least.", f"A good benchmark. Our escape rate is {pct(k['escape_rate'])}."],
       ["Narendranath et al. (2018)", "Used machine learning on the same survey.", "The closest earlier work."],
       ["McBride and Nichols (2018)", "Testing on data the model has not seen, and combining models, improve poverty targeting tools.", "This is why we test on villages the model has not seen."],
       ["Aiken et al. (2022)", "Machine learning targeted cash aid in Togo and left out fewer poor people.", "It shows a real use of this kind of model."]],
      [2.7, 6.2, 3.3], 14, row_h=0.8)
note(s, Inches(0.6), Inches(6.3), Inches(12.2), Inches(0.7),
     "We found no accessible study that predicts poverty transitions of households in India and explains its predictions. That is the gap our project fills.", 14)
notes(s, "Anhad: Earlier work tells us three things. Escaping poverty is real and it differs between groups. The survey we use shows that Adivasi households escaped least. And machine learning is already used to target aid. "
         "Thorat and co-authors used our survey but did not build a prediction model. Narendranath and co-authors is the closest study, but we could not read its results. "
         "So our contribution is a careful prediction, with explanations, tested on villages the model has not seen.")

# =========================================================================================================== 8 Dataset (Aman)
s = new_slide(3, "We use a survey that visited the same households twice")
text(s, Inches(0.6), Inches(1.6), Inches(6.6), Inches(5.2),
     ["**About the survey",
      "The India Human Development Survey (IHDS) was run by the University of Maryland and NCAER. We downloaded it from ICPSR.",
      "**What kind of data",
      "It is panel data: the same households were interviewed in 2004-05 and again in 2011-12. In our data one row is one household, and person records are summarised for each household.",
      "**Linking the two rounds",
      f"The second round does not store the old household number, so we rebuilt it. We found {k['link']['found_again']:,} of the {k['link']['wave1_households']:,} households ({pct(k['link']['share_found_again'], 1)}) again."], 16, gap=0.3)
stat(s, Inches(7.5), Inches(1.6), Inches(2.6), f"{k['n_w1_households']:,}", "households in 2004-05")
stat(s, Inches(10.2), Inches(1.6), Inches(2.6), f"{k['n_w2_households']:,}", "households in 2011-12")
stat(s, Inches(7.5), Inches(3.15), Inches(2.6), f"{k['n_w1_persons']:,}", "people in 2004-05")
stat(s, Inches(10.2), Inches(3.15), Inches(2.6), f"{k['n_kept']:,}", "poor households modelled", ORANGE)
note(s, Inches(7.5), Inches(4.75), Inches(5.3), Inches(1.5), "The data come under the ICPSR terms of use, so we keep them inside the team.", 14)
notes(s, "Aman: Thank you, Anhad. I will now explain the data we used. It comes from the India Human Development Survey, which is a panel: it visited the same households in 2004-05 and again in 2011-12. "
         "The second survey does not store the old household number, so we rebuilt it from the state, district, village and household codes, and we checked it against the official tracking file. "
         "About 83% of the households were found again.")

# =========================================================================================================== 9 Poverty line and target (Aman)
s = new_slide(3, "One poverty line for both rounds defines who escaped")
picture(s, "c13_poverty_line.png", Inches(0.4), Inches(1.5), w=Inches(6.3))
text(s, Inches(6.9), Inches(1.5), Inches(5.9), Inches(5.4),
     ["**The poverty line",
      "A household is poor when its monthly spending per person is below a line. The 2004-05 file uses an older line than the 2011-12 file, so we recomputed 2004-05 with the Tendulkar line used in 2011-12. "
      f"It shows {pct(pr['POOR']['person_weighted'])} of people as poor in 2004-05, close to the published figure of about 37%.",
      "**The target",
      "ESCAPED is 1 if a household was poor in 2004-05 and not poor in 2011-12. It is 0 if the household was poor in both rounds.",
      "**What we found",
      f"Of the {k['n_kept']:,} poor households we followed, {k['n_escaped']:,} escaped ({pct(k['escape_rate'], 1)}) and {k['n_stayed']:,} stayed poor."], 15, gap=0.3)
notes(s, "Aman: Poverty means that spending per person is below a line. The two survey files used different lines, so comparing them directly would not be fair. "
         "The second file also stores the Tendulkar line for 2004-05 prices, so we used that line for both rounds. Our target is 1 if a household was poor in 2004-05 and not poor in 2011-12. "
         "About 69 in every 100 households escaped.")

# =========================================================================================================== 10 Features (Aman)
s = new_slide(3, f"{k['n_features']} features describe each household in 2004-05")
picture(s, "c02_feature_groups.png", Inches(0.4), Inches(1.5), w=Inches(6.4))
text(s, Inches(7.0), Inches(1.5), Inches(5.8), Inches(5.4),
     ["**Where the features come from",
      f"The survey has 945 columns. We chose {k['n_features']} of them using our knowledge of the topic and sorted them into ten groups, such as assets, education, jobs and services. {k['n_numeric']} are numbers and {k['n_nominal']} are categories.",
      "**Features we built",
      "We turned person records into household counts and shares, for example the share of dependants and the number of children under 6.",
      "**What we left out",
      "We left out spending and income, because poverty is defined from spending and they would hand the model the answer. We also left out ID numbers, survey weights and interviewer codes."], 15, gap=0.3)
notes(s, "Aman: Every feature describes the household in 2004-05, so the model only uses what was known before the outcome. We started with 945 survey columns and kept 137 in ten groups, "
         "such as assets, household structure, education, jobs and services. Some of them we built ourselves from the person records. "
         "We left out spending and income on purpose, because poverty is defined from spending and they would give the model the answer.")

# =========================================================================================================== 11 Feature selection (Aman)
full_n = int(fsel.index.max())
small = [n for n in fsel.index if n < full_n and fsel.loc[n, "Test ROC-AUC"] >= fsel.loc[full_n, "Test ROC-AUC"] - 0.01]
if small:
    n_small = int(min(small))
    title11 = f"{n_small} features work almost as well as all {full_n}"
    verdict = (f"With only the top {n_small} features the test ROC-AUC is {fsel.loc[n_small, 'Test ROC-AUC']:.3f}, close to {fsel.loc[full_n, 'Test ROC-AUC']:.3f} with all {full_n}. "
               "A much smaller set works almost as well.")
else:
    title11 = "Removing features lowers the score, so we keep all of them"
    verdict = f"The score falls when features are removed ({fsel.loc[full_n, 'Test ROC-AUC']:.3f} with all {full_n}), so we keep all of them."
s = new_slide(3, title11)
picture(s, "c16_feature_selection.png", Inches(0.4), Inches(1.5), w=Inches(6.6))
text(s, Inches(7.2), Inches(1.5), Inches(5.6), Inches(5.4),
     ["**What we did",
      "We ranked the features by how much an XGBoost model used them, looking at the training villages only. Then we kept the top 10, 20, 40 or 80 and trained XGBoost again.",
      "**What we found", verdict,
      "**What we decided",
      f"We keep all {full_n} features in the main model, because nothing is lost and the explanations then cover every group of factors. The smaller set shows that the model does not depend on a long list of weak features."], 15, gap=0.3)
notes(s, "Aman: The coordinator's email asks about feature selection, so we tested it directly. We ranked the features using only the training villages, kept the top few, and trained the model again. "
         "The chart shows how the score changes. With about 40 features the score is already close to the score with all of them. We still keep all 137, for the reason on the slide.")

# =========================================================================================================== 12 Missing values and preprocessing (Aman)
s = new_slide(3, "Missing values and preprocessing")
miss = dictionary[dictionary["used"]].sort_values("missing_share", ascending=False).head(6)
miss = miss.head(5)
rows = [["Feature with most gaps", "Share missing"]] + [[str(lbl)[:48], pct(m)] for lbl, m in zip(miss["label"], miss["missing_share"])]
table(s, Inches(0.6), Inches(1.5), rows, [4.6, 1.6], 14, row_h=0.5)
text(s, Inches(0.6), Inches(5.2), Inches(6.2), Inches(1.6),
     ["The survey writes unanswered questions as codes from -1 to -9. We turned these into blanks. One feature with more than 40% blanks was dropped."], 14)
text(s, Inches(7.1), Inches(1.5), Inches(5.7), Inches(5.4),
     ["**Four steps, all learned from the training villages only",
      "1. Blank numbers get the middle value (the median), and blank categories get their own 'missing' category.",
      "2. Extreme values are pulled in to the 0.5th and 99.5th percentile, so a typing error cannot distort the model.",
      "3. Categories such as state become yes/no columns (one-hot encoding).",
      "4. Numbers are scaled for Logistic Regression only.",
      "**One feature we treat with care",
      "The ration card type is a label that the government gives to poor households, so we do not read it as a risk factor."], 15, gap=0.28)
notes(s, "Aman: Missing answers were stored as negative codes, so we turned them into blanks. Blank numbers get the median, and extreme values are pulled in, because a few cells hold typing errors. "
         "Categories such as state become yes or no columns. All of this is learned from the training villages only, so nothing leaks from the test villages.")

# =========================================================================================================== 13 EDA (Aman)
s = new_slide(3, "Escape rates differ by social group and by education")
picture(s, "c03_escape_by_caste.png", Inches(0.3), Inches(1.5), w=Inches(6.4))
picture(s, "c04_escape_by_education.png", Inches(6.7), Inches(1.5), w=Inches(6.4))
caste = chart_data.groupby("Caste/religion", observed=True)["ESCAPED"].agg(["mean", "size"])
educ = chart_data.groupby("Adult education", observed=True)["ESCAPED"].agg(["mean", "size"])
text(s, Inches(0.6), Inches(5.55), Inches(6.0), Inches(1.4),
     [f"Adivasi households escaped least often ({pct(caste.loc['Adivasi', 'mean'])}). Dalit households ({pct(caste.loc['Dalit', 'mean'])}) were close to the overall average of {pct(k['escape_rate'])}."], 15)
text(s, Inches(7.0), Inches(5.55), Inches(5.8), Inches(1.4),
     [f"The escape rate rises with schooling: {pct(educ.loc['None', 'mean'])} when no adult had any schooling and {pct(educ.loc['13+ yrs', 'mean'])} when an adult had 13 or more years. These are raw differences."], 15)
notes(s, "Aman: Before building any model we looked at the data. Adivasi households escaped least often, and the escape rate rises with the schooling of the most educated adult. "
         "These are raw comparisons that ignore other differences between households, which is one reason we use models.")

# =========================================================================================================== 14 Validation (Eshaan)
s = new_slide(4, "We test on whole villages that the models have never seen")
stat(s, Inches(0.6), Inches(1.6), Inches(2.9), f"{k['n_train']:,}", "training households")
stat(s, Inches(3.7), Inches(1.6), Inches(2.9), f"{k['n_train_villages']:,}", "training villages")
stat(s, Inches(0.6), Inches(3.15), Inches(2.9), f"{k['n_test']:,}", "test households")
stat(s, Inches(3.7), Inches(3.15), Inches(2.9), f"{k['n_test_villages']:,}", "test villages")
box(s, Inches(0.6), Inches(4.7), Inches(6.0), Inches(2.1), "Why split by village?",
    ["Households in one village are alike and share roads, prices and weather. If neighbours sit in both the training part and the test part, the test score looks better than it really is. (In our data a village can also be an urban neighbourhood.)"], 14)
if rv["random_mean"] >= rv["village_mean"]:
    cmp_line = f"We repeated both kinds of split 20 times. A random split of households gives an average ROC-AUC of {rv['random_mean']:.3f} and a split by village gives {rv['village_mean']:.3f}. The random split looks better because neighbours are in both parts."
else:
    cmp_line = (f"We repeated both kinds of split 20 times. A random split of households gives an average ROC-AUC of {rv['random_mean']:.3f} and a split by village gives {rv['village_mean']:.3f}. "
                "The two are close here, but only the village split guarantees that neighbours never sit on both sides.")
box(s, Inches(6.9), Inches(1.6), Inches(5.9), Inches(3.0), "Random split compared with village split", [cmp_line], 15)
box(s, Inches(6.9), Inches(4.8), Inches(5.9), Inches(2.0), "Cross-validation",
    ["While tuning, we cut the training villages into five groups, train on four, check on the fifth, and rotate. A village is never split between groups."], 14)
notes(s, "Eshaan: Thank you, Aman. I will now explain how we set up and test the models. We split the data by village and not by household. "
         "Households in one village share roads, prices and weather, so a random split would put neighbours on both sides and make the score look better than it is. "
         "We checked this by repeating both kinds of split 20 times. The test villages are never used to choose a model or its settings.")

# =========================================================================================================== 15 Workflow (Eshaan)
s = new_slide(4, "The machine learning pipeline has nine steps")
picture(s, "c17_workflow.png", Inches(0.4), Inches(1.45), w=Inches(12.5))
text(s, Inches(0.6), Inches(6.3), Inches(12.2), Inches(0.7),
     ["The test villages are not touched until step 7, when the finished models are scored on them for the first time. Everything before that, including the tuning in step 6, uses the training villages only."], 14)
notes(s, "Eshaan: Here is the whole pipeline in nine boxes. We check and link the survey files, measure poverty with one line, create the target, build the features and split by village. "
         "Then we train and tune four models, test them on villages they have never seen, explain the best one, and save everything. The slides, the report and the app all read the same saved results.")

# =========================================================================================================== 16 Algorithms (Eshaan)
s = new_slide(4, "Algorithms and why we chose them")
table(s, Inches(0.6), Inches(1.5),
      [["Model", "How it works", "Why we chose it"],
       ["Dummy", "Always predicts the overall escape rate", "Shows what having no skill looks like"],
       ["Logistic Regression", "Adds up weighted facts and turns the total into a probability", "Simple and easy to explain"],
       ["Random Forest", "Averages the answers of many decision trees", "Can find curved patterns and interactions"],
       ["XGBoost", "Builds trees one after another, each fixing the errors of the last", "Works well on tables of data like ours"],
       ["Ensemble", "Averages the probabilities of the three models", "Combines a straight-line model with tree models"]],
      [2.6, 5.1, 4.5], 15, row_h=0.62)
note(s, Inches(0.6), Inches(5.6), Inches(12.2), Inches(1.2),
     "Why not a neural network or an SVM? We have only about 9,000 rows of mixed survey data. For this kind of table, tree models are the usual strong choice, and they are easier to explain.", 15)
notes(s, "Eshaan: We chose the models from simple to strong. Logistic Regression is the simple baseline. Random Forest and XGBoost can find curved patterns and interactions. "
         "The ensemble averages the three, and the dummy shows what no skill looks like. We did not use a neural network because we have only about nine thousand rows, and trees are easier to explain.")

# =========================================================================================================== 17 Training strategy (Eshaan)
s = new_slide(4, "Training strategy")
box(s, Inches(0.6), Inches(1.6), Inches(5.9), Inches(2.5), "Tuning",
    ["Every model has settings that must be chosen before training. A random search tries many settings and keeps the one with the best cross-validated ROC-AUC."], 15)
box(s, Inches(6.8), Inches(1.6), Inches(6.0), Inches(2.5), "Class balance",
    [f"About {pct(k['escape_rate'])} of the households escaped, so the two classes are a little unbalanced. We did not use class weights, because they distort the predicted probabilities."], 15)
box(s, Inches(0.6), Inches(4.3), Inches(5.9), Inches(2.5), "Cut-off",
    [f"A cut-off turns a probability into a decision, 'escaped' or 'stayed poor'. For XGBoost it is {thr[HEAD]:.3f}. We chose it on training predictions, never on the test set."], 15)
box(s, Inches(6.8), Inches(4.3), Inches(6.0), Inches(2.5), "Headline model",
    [f"XGBoost has the best cross-validated ROC-AUC ({cv_summary.loc['XGBoost', 'CV ROC-AUC (mean)']:.3f}), so it is our headline model. We made this choice before we looked at the test set."], 15)
notes(s, "Eshaan: Tuning means trying many settings and keeping the best one according to cross-validation. We did not use class weights, because they make the predicted probabilities wrong. "
         "The cut-off is picked on training predictions so that the average F1 of both classes is as high as possible. And we chose XGBoost as the headline model by cross-validation, before we looked at the test set.")

# =========================================================================================================== 18 Implementation and structure (Aman)
s = new_slide(5, "Implementation and project structure")
mono(s, Inches(0.6), Inches(1.5), Inches(7.4), Inches(5.3), [
    "AUCML project folder",
    "|-- IHDS_poverty_escape_v3.ipynb   main notebook",
    "|-- ihds_tools.py                  helper class",
    "|-- requirements.txt               package versions",
    "|-- quick_check.py                 re-check results",
    "|-- app.py, app_labels.json        Streamlit demo",
    "|-- make_charts_v3.py              charts",
    "|-- build_deck_v4.py               slides",
    "|-- outputs/",
    "|     |-- *.csv, key_numbers.json  tables and numbers",
    "|     |-- figures/, figures_ppt/   charts",
    "|     |-- model_bundle.joblib      saved models",
    "|-- AUCML_Project_Report.docx/pdf  report",
    "|-- ICPSR_22626-V12/, ICPSR_36151-V6/",
    "                                   raw data (not shared)"], 13)
text(s, Inches(8.3), Inches(1.5), Inches(4.5), Inches(5.3),
     ["**Tools",
      "Python 3.11 with pandas, NumPy, scikit-learn, XGBoost, SHAP and matplotlib, and Streamlit for the demo.",
      "**How it works",
      "The notebook is run from top to bottom and writes all tables, charts and saved models into the outputs folder. The charts, slides, report and app then read from that folder.",
      "**Why the results repeat",
      "We fixed the random seeds, saved the cross-validation folds and pinned the package versions."], 14, gap=0.3)
notes(s, "Aman: Thank you, Eshaan. I will now show how the project is built and run. This is our project folder, and the notebook is the heart of it. You run it from top to bottom, "
         "and it writes the tables, charts and saved models into the outputs folder. Everything else, meaning the charts, the slides, the report and the app, reads from outputs. "
         "We fixed the random seeds and pinned the package versions so that the results repeat.")

# =========================================================================================================== 19 Execution evidence (Aman)
s = new_slide(5, "The code runs, and the results can be re-checked in seconds")
picture(s, "exec_evidence.png", Inches(0.45), Inches(1.4), w=Inches(6.3))
mono(s, Inches(7.0), Inches(1.45), Inches(5.85), Inches(3.45), QUICK_CHECK_LINES, 10.5)
text(s, Inches(0.6), Inches(5.4), Inches(12.2), Inches(1.6),
     ["The picture on the left is the real output of the executed notebook, with the tuning time of each model. The box on the right is the real output of quick_check.py, which recomputes the headline scores from the saved predictions in seconds.",
      "The full notebook takes about 30 minutes on our laptop. Running it again gives identical result tables, even with another scikit-learn version."], 14, gap=0.3)
notes(s, "Aman: This slide shows that the project was really executed. On the left is the output of the notebook run, with the time it took to tune each model. "
         "On the right is the output of our quick check script, which recomputes the headline scores from the saved predictions in a few seconds, so we can run it live if you like. "
         "The full notebook takes about 30 minutes, and running it again gave identical result tables.")

# =========================================================================================================== 20 Metrics (Gokul)
s = new_slide(6, "Evaluation metrics and what they mean")
table(s, Inches(0.6), Inches(1.5),
      [["Metric", "In plain words"],
       ["Accuracy", "The share of all predictions that are right"],
       ["Precision", "Of the households we call 'escaped', the share that really escaped"],
       ["Recall", "Of the households that really escaped, the share we found"],
       ["F1-score", "One number that balances precision and recall"],
       ["Confusion matrix", "A table of right and wrong predictions"],
       ["Macro-F1", "The F1-score averaged over both classes"],
       ["ROC-AUC", "The chance that the model scores a random escaped household above a random stayed-poor one. 0.5 is guessing and 1 is perfect"],
       ["95% interval", "How far a score could move by chance"]],
      [2.6, 9.6], 14, row_h=0.52)
note(s, Inches(0.6), Inches(6.35), Inches(12.2), Inches(0.6),
     f"Accuracy alone can mislead here. A model that always says 'escaped' would be right {pct(res['Dummy (predict prior)']['Accuracy'])} of the time without learning anything.", 15)
notes(s, "Gokul: Thank you, Aman. I will now take you through the results. First, the measures we use. Accuracy is the share of correct answers, but it can mislead here, "
         "because a model that always says escaped is right about 69% of the time. So we also report precision, recall and F1 for both groups, the confusion matrix, "
         "and ROC-AUC, which does not depend on the cut-off.")

# =========================================================================================================== 21 Comparison summary (Gokul)
s = new_slide(6, "XGBoost scores highest, but all four models are close")
order = [HEAD, ENS, "Random Forest", "Logistic Regression", "Dummy (predict prior)"]
names = {ENS: "Ensemble", "Dummy (predict prior)": "Dummy"}
rows = [["Model", "ROC-AUC (95% interval)", "Accuracy", "Macro-F1"]]
for m in order:
    r = res[m]
    ci = "" if m.startswith("Dummy") else " (" + r["ROC_AUC_CI"].strip("[]").replace(", ", " to ") + ")"
    rows.append([names.get(m, m), f"{r['ROC_AUC']:.3f}{ci}", f"{r['Accuracy']:.3f}", f"{r['Macro_F1']:.3f}"])
table(s, Inches(0.6), Inches(1.5), rows, [3.0, 4.6, 2.2, 2.4], 16, highlight=1, row_h=0.6)
pe = pair[(ENS, HEAD)]
box(s, Inches(0.6), Inches(5.3), Inches(12.2), Inches(1.55), "Reading the table",
    [f"XGBoost and the ensemble score highest, and the gap between them ({abs(pe['difference']):.3f}) is within chance. All four models beat the dummy, but the scores are modest, between 0.70 and 0.72. "
     "The dummy gives every household the same score, so its true ROC-AUC is 0.5, and the 0.487 comes from tiny random noise that breaks ties."], 14)
notes(s, "Gokul: Here are the four models on the same test villages. XGBoost scores 0.722, the ensemble 0.719, Random Forest 0.709 and Logistic Regression 0.700, and the dummy scores 0.487, which is guessing. "
         "The top two are within chance of each other. The scores are modest, which is why we call this a ranking tool.")

# =========================================================================================================== 22 Per-class metrics (Gokul)
s = new_slide(6, "All four models find escaped households more easily")
rows = [["Model", "Prec. escaped", "Recall escaped", "F1 escaped", "Prec. stayed poor", "Recall stayed poor", "F1 stayed poor"]]
for m in ["Logistic Regression", "Random Forest", "XGBoost", ENS]:
    r = fm[m]
    rows.append([names.get(m, m)] + [f"{r[c]:.3f}" for c in ["Precision (escaped)", "Recall (escaped)", "F1 (escaped)", "Precision (stayed poor)", "Recall (stayed poor)", "F1 (stayed poor)"]])
table(s, Inches(0.6), Inches(1.5), rows, [2.3, 1.65, 1.65, 1.5, 1.8, 1.8, 1.6], 15, highlight=3, row_h=0.65)
text(s, Inches(0.6), Inches(5.0), Inches(12.2), Inches(1.8),
     ["Every model finds the households that escaped fairly well (recall about 0.73 to 0.77), and finds the households that stayed poor less well (recall about 0.53 to 0.57).",
      "A welfare programme cares most about the households that stay poor, so we report both groups."], 15, gap=0.3)
notes(s, "Gokul: This table gives precision, recall and F1 for each group. All the models show the same pattern: they find the escaped households better than the stayed-poor households. "
         "A welfare programme cares most about the stayed-poor households, which is why we show both groups.")

# =========================================================================================================== 23 ROC (Gokul)
s = new_slide(6, "The four ROC curves are close, and XGBoost is slightly ahead")
picture(s, "c05_roc_curves.png", Inches(0.5), Inches(1.5), h=Inches(5.3))
text(s, Inches(6.9), Inches(1.6), Inches(5.9), Inches(5.3),
     ["**How to read it",
      "The further a curve rises above the dashed line, which is guessing, the better the model ranks households. Here the four curves are close together.",
      "**Cross-validation tells the same story",
      f"The average ROC-AUC on the validation folds of the training villages was {cv_summary.loc['Logistic Regression', 'CV ROC-AUC (mean)']:.3f} for Logistic Regression, {cv_summary.loc['Random Forest', 'CV ROC-AUC (mean)']:.3f} for Random Forest and {cv_summary.loc['XGBoost', 'CV ROC-AUC (mean)']:.3f} for XGBoost. "
      "XGBoost is the best, which is why it is our headline model."], 15, gap=0.3)
notes(s, "Gokul: The ROC curve shows how well a model ranks households at every possible cut-off. All four curves are well above the dashed line, which is guessing, and they are close to each other. "
         "Cross-validation on the training villages tells the same story and points to XGBoost.")

# =========================================================================================================== 24 Confusion matrices (Gokul)
s = new_slide(6, "XGBoost finds just over half of the households that stayed poor")
picture(s, "c18_confusion_all_models.png", Inches(0.3), Inches(1.5), w=Inches(12.7))
cmx = fm[HEAD]
text(s, Inches(0.6), Inches(5.2), Inches(12.2), Inches(1.7),
     [f"Take the XGBoost matrix as an example. Of the {cmx['True stayed'] + cmx['False escaped']:,} households that stayed poor, the model found {cmx['True stayed']:,}, and of the {cmx['True escaped'] + cmx['False stayed']:,} that escaped, it found {cmx['True escaped']:,}.",
      "The model misses many households, so we use it to rank households and not to decide about one family."], 15, gap=0.3)
notes(s, "Gokul: Each box compares what a model predicted with what really happened. For XGBoost, the model finds a little over half of the households that stayed poor and about three quarters of the households that escaped. "
         "All four models look alike. They make many errors, so we use the model to rank households and not to decide about one family.")

# =========================================================================================================== 25 Risk groups (Gokul)
s = new_slide(6, f"The riskiest 20% of households stayed poor {pct(lift[-1]['stayed_poor_rate'])} of the time")
picture(s, "c07_risk_groups.png", Inches(0.5), Inches(1.5), w=Inches(7.4))
top40 = lift[-1]["share_of_all_stayed_poor"] + lift[-2]["share_of_all_stayed_poor"]
text(s, Inches(8.2), Inches(1.6), Inches(4.6), Inches(5.3),
     ["**What we did",
      "We sorted the test households by predicted risk and cut them into five equal groups.",
      "**What we found",
      f"In the highest-risk group {pct(lift[-1]['stayed_poor_rate'])} stayed poor, which is {lift[-1]['lift_vs_average']:.1f} times the average of {pct(1 - float(np.mean(y_test)))}. In the lowest-risk group only {pct(lift[0]['stayed_poor_rate'])} did.",
      f"The riskiest 40% of households contain {pct(top40)} of all the households that stayed poor."], 15, gap=0.3)
notes(s, "Gokul: This is the practical value of the model. We ranked the test households by predicted risk and cut them into five equal groups. In the highest-risk group, 57% stayed poor, "
         "against 31% on average and only 11% in the lowest group. A programme with a small budget could focus on the top groups and still reach more than half of the households that stay poor.")

# =========================================================================================================== 26 SHAP top (Eshaan)
s = new_slide(7, "Young children, state and household assets matter most")
picture(s, "c08_shap_top10.png", Inches(0.4), Inches(1.5), w=Inches(7.6))
text(s, Inches(8.3), Inches(1.6), Inches(4.5), Inches(5.3),
     ["**What SHAP does",
      "SHAP splits a prediction into the push that each feature gave it, a little like splitting a bill between items. We average these pushes over many households.",
      "**What it shows",
      "The model relies most on the number of young children, the state, household assets, the ration card type, and the caste and religion group.",
      "**What to keep in mind",
      "SHAP shows what the model uses. It does not show what causes escape."], 14, gap=0.3)
notes(s, "Eshaan: Thank you, Gokul. Next I will explain what the model has learned. SHAP splits each prediction into the push from each feature, a little like splitting a bill between items. "
         "Averaging these pushes shows what the model relies on: the number of young children, the state, assets, the ration card type and the caste and religion group. "
         "This tells us what the model uses, and it does not tell us what causes escape. The ration card is a label the government gives to poor households, so we do not read it as a risk factor.")

# =========================================================================================================== 27 SHAP direction (Eshaan)
s = new_slide(7, "Fewer young children and more assets point towards escape")
picture(s, "c10_shap_children.png", Inches(0.8), Inches(1.5), w=Inches(5.1))
picture(s, "c11_shap_assets.png", Inches(7.1), Inches(1.5), w=Inches(5.1))
text(s, Inches(0.6), Inches(5.55), Inches(5.9), Inches(1.4),
     ["Households with no child under 6 are pushed towards escape, and households with two or more are pushed towards staying poor."], 15)
text(s, Inches(6.9), Inches(5.55), Inches(5.9), Inches(1.4),
     ["Households with more assets per person are generally pushed towards escape."], 15)
notes(s, "Eshaan: Each dot is a household. A dot above zero means the feature pushed the prediction towards escape. Households with no child under six are pushed towards escape, "
         "and households with two or more are pushed towards staying poor. More assets per person generally push towards escape.")

# =========================================================================================================== 28 Education (Eshaan)
s = new_slide(7, "Households with an educated adult escaped more often")
picture(s, "c19_education_simple.png", Inches(0.4), Inches(1.5), w=Inches(6.6))
text(s, Inches(7.2), Inches(1.5), Inches(5.6), Inches(5.4),
     ["**The question",
      f"Among poor households, do those with an adult who has 10 or more years of schooling ({edu['n_treated']:,} households) escape more often?",
      "**The answer",
      f"Yes. They escaped {EA['raw']['estimate'] * 100:.1f} percentage points more often. After comparing them with similar households (same state, caste, area, family structure and land), the gap is about {EA['doubly_robust']['estimate'] * 100:.0f} points.",
      "**What to keep in mind",
      "This is a link between education and escape. It does not prove that education causes escape."], 15, gap=0.3)
notes(s, "Eshaan: Our registered idea included a small propensity-score analysis, and this is it. Poor households with an educated adult escaped about 11 points more often. "
         "After comparing them with similar households, using weights built from propensity scores, the gap is about 6 points. We still call this a link and not a cause. "
         "The details and the safety checks are in the backup slides.")

# =========================================================================================================== 29 Limitations (Anhad)
s = new_slide(8, "The model is modest, and the data have limits")
items = [("Modest accuracy", "A ROC-AUC of about 0.7 ranks households reasonably well, but many single predictions are wrong."),
         ("A noisy poverty line", "Escaping is partly a matter of crossing a line. One in five poor households were within 10% of it."),
         ("Households we lost", f"About {pct(1 - k['retention'])} of poor households could not be followed up, for example because they moved."),
         ("Old data", "The survey covers 2004 to 2012, so newer data is needed before any real use."),
         ("Patterns, not causes", "SHAP and the education result describe patterns. They do not show what causes escape."),
         ("Fairness", "State and caste are strong predictors, so the score should support groups and never be used to leave them out.")]
for i, (h, b) in enumerate(items):
    x = Inches(0.6) + Inches(6.15) * (i % 2)
    y = Inches(1.5) + Inches(1.8) * (i // 2)
    box(s, x, y, Inches(6.0), Inches(1.65), h, [b], 14)
notes(s, "Anhad: Thank you, Eshaan. We want to be honest about the limits. The model can rank households, but it is not accurate enough to judge a single family. "
         "Escaping is partly noise around a poverty line, and about 11% of poor households were lost between the surveys. The data are old. Our results show patterns and not causes. "
         "And because state and caste are strong predictors, the score must be used to support groups and never to leave them out.")

# =========================================================================================================== 30 Conclusion (Anhad)
s = new_slide(8, "Conclusion and future scope")
box(s, Inches(0.6), Inches(1.5), Inches(7.6), Inches(5.3), "Conclusion",
    [f"1. About 7 in 10 poor households in 2004-05 had escaped by 2011-12 ({pct(k['escape_rate'], 1)}).",
     f"2. Our best model, XGBoost, scores a ROC-AUC of {res[HEAD]['ROC_AUC']:.3f} on villages it has never seen. That is a modest score.",
     f"3. The 20% of households rated riskiest stayed poor {pct(lift[-1]['stayed_poor_rate'])} of the time, against {pct(1 - float(np.mean(y_test)))} on average.",
     "4. Young children, state, assets, the ration card and the caste and religion group matter most to the model.",
     "5. Adult education goes together with a 5 to 6 point higher escape rate, but this does not prove that education is the cause."], 15, gap=0.45)
box(s, Inches(8.4), Inches(1.5), Inches(4.4), Inches(5.3), "Future scope",
    ["Retrain on newer panel data.", "Add satellite or village data.", "Predict falling into poverty.", "Check fairness for each group.", "Extend the app for field workers."], 15, gap=0.5)
notes(s, "Anhad: To conclude, about seven in ten poor households escaped. Our best model ranks households modestly well, and the riskiest fifth stayed poor 57% of the time. "
         "The strongest signals are young children, state, assets, the ration card and social group. Education goes together with escape, but we do not claim that it causes it. "
         "In future we would use newer data, add village data, build a model for falling into poverty, and check fairness for each group.")

# =========================================================================================================== 31 Demo (Anhad)
s = new_slide(8, "Demo: the poverty risk app")
picture(s, "app_screenshot.png", Inches(0.5), Inches(1.45), w=Inches(7.9))
text(s, Inches(8.6), Inches(1.5), Inches(4.2), Inches(5.4),
     ["**Step 1", "Enter a household's 2004-05 details, or start from a real test household.",
      "**Step 2", "See the chance that the household escapes, and its risk group.",
      "**Step 3", "See which facts raised or lowered the estimate (SHAP).",
      "**A caution", "The app is a screening aid. It should not be used to decide about one household."], 14, gap=0.25)
notes(s, "Anhad: Now a short live demo. We run streamlit run app.py and choose the real test household with the lowest predicted chance of escaping. "
         "The app shows the chance of escaping, the risk group, how often households in that group really stayed poor, and a chart of what pushed the estimate. "
         "Then I change the number of assets so you can see the estimate move. Please remember that it is a screening aid. Thank you, and we are happy to take your questions.")

# =========================================================================================================== 32 References
s = new_slide(None, "References")
refs = [
    "1. Krishna, A. (2004). Escaping poverty and becoming poor: Who gains, who loses, and why? World Development, 32(1), 121-136. https://doi.org/10.1016/j.worlddev.2003.08.002",
    "2. Thorat, A., Vanneman, R., Desai, S., and Dubey, A. (2017). Escaping and falling into poverty in India today. World Development, 93, 413-426. https://doi.org/10.1016/j.worlddev.2017.01.004",
    "3. Narendranath, S., Khare, S., Gupta, D., and Jyotishi, A. (2018). Characteristics of Escaping and Falling into Poverty in India: An Analysis of IHDS Panel Data using machine learning approach. ICACCI 2018, 1391-1397. https://doi.org/10.1109/ICACCI.2018.8554571",
    "4. McBride, L., and Nichols, A. (2018). Retooling poverty targeting using out-of-sample validation and machine learning. World Bank Economic Review, 32(3), 531-550. https://doi.org/10.1093/wber/lhw056",
    "5. Aiken, E., Bellue, S., Karlan, D., Udry, C., and Blumenstock, J. E. (2022). Machine learning and phone data can improve targeting of humanitarian aid. Nature, 603, 864-870. https://doi.org/10.1038/s41586-022-04484-9",
    "6. Jean, N., Burke, M., Xie, M., Davis, W. M., Lobell, D. B., and Ermon, S. (2016). Combining satellite imagery and machine learning to predict poverty. Science, 353(6301), 790-794. https://doi.org/10.1126/science.aaf7894",
    "7. Breiman, L. (2001). Random forests. Machine Learning, 45, 5-32. https://doi.org/10.1023/A:1010933404324",
    "8. Chen, T., and Guestrin, C. (2016). XGBoost: A scalable tree boosting system. KDD 2016, 785-794. https://doi.org/10.1145/2939672.2939785",
    "9. Lundberg, S. M., and Lee, S.-I. (2017). A unified approach to interpreting model predictions. NeurIPS 30. https://arxiv.org/abs/1705.07874",
    "10. VanderWeele, T. J., and Ding, P. (2017). Sensitivity analysis in observational research: introducing the E-value. Annals of Internal Medicine, 167(4), 268-274. https://doi.org/10.7326/M16-2607",
    "11. Planning Commission (2009). Report of the Expert Group to Review the Methodology for Estimation of Poverty (Tendulkar Committee). Government of India.",
    "12. Desai, S., Vanneman, R., and NCAER. India Human Development Survey (IHDS), ICPSR 22626 (2004-05) and IHDS-II, ICPSR 36151 (2011-12). https://ihds.umd.edu",
    "13. United Nations. Sustainable Development Goals 1 and 10. https://sdgs.un.org/goals/goal1 and https://sdgs.un.org/goals/goal10",
]
text(s, Inches(0.6), Inches(1.3), Inches(12.2), Inches(5.7), refs, 10.5, gap=0.3)
notes(s, "References. All are checked: the DOIs resolve to the titles and pages shown.")

# =========================================================================================================== 33 Backup divider
s = new_slide(None, "Backup slides")
rect(s, Inches(0.6), Inches(1.6), Inches(12.2), Inches(4.8), fill=BACKUP_FILL, line=None)
text(s, Inches(1.0), Inches(2.0), Inches(11.4), Inches(4.2),
     ["**The slides after this one are for questions only",
      "They hold the extra checks, such as how we linked the rounds and the households we lost, calibration, how stable the ranking is, robustness checks, "
      "the education analysis in detail, SHAP by group and a short glossary.",
      "We do not present them. We use them if an examiner asks."], 20, gap=0.5)
notes(s, "Backup divider. The slides after this one are only for questions.")

# =========================================================================================================== B1 Linking and attrition
s = backup_slide("Linking the rounds and the households we lost")
picture(s, "c01_transitions.png", Inches(0.4), Inches(1.5), w=Inches(6.4))
text(s, Inches(7.0), Inches(1.5), Inches(5.8), Inches(5.4),
     ["**The link",
      f"We rebuilt it from state, district, village, household and split codes. Following the 2005 head of household instead gives the same outcome for {pct(k['head_rule_check']['same_outcome'], 1)} of the {pct(k['head_rule_check']['share_head_tracked'])} of households where the head can be tracked.",
      "**The households we lost",
      f"{pct(1 - k['retention'])} of poor households could not be followed. They are more often urban and own less land. When we give extra weight to the households we did follow, the escape rate moves from {pct(k['ipw']['escape_rate_raw'], 1)} to {pct(k['ipw']['escape_rate_ipw'], 1)}."], 14, gap=0.3)
notes(s, "If asked: The link is built from the survey codes and agrees with the official tracking file. About 11% of poor households were lost, often because they moved. Re-weighting for this changes the escape rate very little.")

# =========================================================================================================== B2 Calibration
s = backup_slide("Can the predicted probabilities be trusted?")
picture(s, "c14_calibration.png", Inches(0.5), Inches(1.5), h=Inches(5.0))
cx = cal[HEAD]
text(s, Inches(6.6), Inches(1.6), Inches(6.2), Inches(5.3),
     ["**What it shows",
      "We group households by their predicted chance of escaping and compare it with the share that really escaped. Points on the dashed line mean the probability is exactly right.",
      "**What we found",
      f"Average predicted chance {pct(cx['mean_predicted'], 1)}, real rate {pct(cx['observed_rate'], 1)}. Brier score {cx['Brier']:.3f} (average squared error of the probabilities, lower is better). "
      "So a predicted 70% means that roughly 70 in 100 similar households escaped.",
      "We did not use class weights because they pull the predicted chances away from the real rates."], 14, gap=0.3)
notes(s, "If asked: Calibration asks whether a predicted 70% really means about 70 in 100. Ours is close. We avoided class weights because they distort the probabilities.")

# =========================================================================================================== B3 Stability and overfitting
s = backup_slide("Stability of the ranking and overfitting check")
picture(s, "c15_repeated_splits.png", Inches(0.4), Inches(1.5), w=Inches(6.6))
rows = [["Model", "Train", "Cross-val.", "Test"]]
for m in ["Logistic Regression", "Random Forest", "XGBoost"]:
    rows.append([m, f"{over[m]['Train ROC-AUC']:.3f}", f"{over[m]['Cross-validated ROC-AUC']:.3f}", f"{over[m]['Test ROC-AUC']:.3f}"])
table(s, Inches(7.2), Inches(1.5), rows, [2.4, 1.0, 1.2, 1.0], 13, row_h=0.5)
text(s, Inches(7.2), Inches(3.8), Inches(5.6), Inches(3.1),
     ["The chart on the left shows the test ROC-AUC over 20 different village splits. The ensemble and XGBoost lead and are very close, and one split can move a score by about 0.015.",
      "The table above shows that Random Forest fits the households it learned from almost perfectly (0.988), so we judge models on cross-validation and test villages only."], 14, gap=0.3)
notes(s, "If asked: We repeated the train and test step on 20 village splits. The order of models is stable, but the top two swap places sometimes. The table shows overfitting: Random Forest fits its training households almost perfectly, "
         "so training scores are not used to judge models.")

# =========================================================================================================== B4 Robustness
s = backup_slide("Robustness checks")
f_final, f_village, f_mates, f_spend = k["feature_sets"]
sens = {r["Target rule"]: r for r in k["sensitivity"]}
old_line = {r["Model"]: r for r in k["old_line"]}
sw = k["survey_weights"]
rows = [["Check", "Result on test villages"],
        ["Add village infrastructure data", f"ROC-AUC: LR {f_village['LR test AUC']:.3f}, XGBoost {f_village['XGB test AUC']:.3f}"],
        ["Add average traits of other households in the village", f"ROC-AUC: LR {f_mates['LR test AUC']:.3f}, XGBoost {f_mates['XGB test AUC']:.3f}"],
        ["Add back spending and income", f"ROC-AUC: LR {f_spend['LR test AUC']:.3f}, XGBoost {f_spend['XGB test AUC']:.3f}"],
        ["Escape only if all split-off families are non-poor", f"Escape rate {pct(sens['Escaped only if ALL descendant households non-poor']['Escape rate'], 1)}, ROC-AUC {sens['Escaped only if ALL descendant households non-poor']['ROC_AUC (test)']:.3f}"],
        ["Escape if any split-off family is non-poor", f"Escape rate {pct(sens['Escaped if ANY descendant household non-poor']['Escape rate'], 1)}, ROC-AUC {sens['Escaped if ANY descendant household non-poor']['ROC_AUC (test)']:.3f}"],
        ["Only households also poor on the older line", f"{int(old_line['XGBoost']['households (deeper poor)']):,} households, ROC-AUC (XGBoost) {old_line['XGBoost']['AUC: main model']:.3f}"],
        ["Survey weights", f"Weighted escape rate {pct(sw['escape_rate_weighted'], 1)}, weighted ROC-AUC (XGBoost) {[a for a in sw['auc'] if a['Model'] == 'XGBoost'][0]['AUC weighted']:.3f}"],
        ["Without the ration card", f"ROC-AUC (XGBoost) {k['ration_card'][1]['AUC without ration card']:.3f} instead of {k['ration_card'][1]['AUC with ration card']:.3f}"]]
table(s, Inches(0.6), Inches(1.5), rows, [6.2, 6.0], 13, row_h=0.5)
text(s, Inches(0.6), Inches(6.2), Inches(12.2), Inches(0.8),
     ["Adding spending and income back changes the score by less than 0.01, so leaving them out costs almost nothing."], 14)
notes(s, "If asked: These checks change one thing at a time and see whether the main result moves. It hardly does. Leaving out spending and income, which we did on purpose, costs very little.")

# =========================================================================================================== B5 Education in detail
s = backup_slide("Education analysis in detail")
picture(s, "c12_education_effect.png", Inches(0.4), Inches(1.5), w=Inches(7.6))
text(s, Inches(8.2), Inches(1.5), Inches(4.6), Inches(5.4),
     ["**Three ways of comparing similar households",
      f"Regression, weighting by propensity scores and a mix of the two (called doubly robust) give {EA['regression']['estimate'] * 100:.1f}, {EA['ipw']['estimate'] * 100:.1f} and {EA['doubly_robust']['estimate'] * 100:.1f} points. The intervals come from re-drawing whole villages.",
      "**The E-value",
      f"The E-value is {EA['doubly_robust']['e_value']:.1f}. A hidden factor that raised both education and escape by that much could explain the gap away, so the result is only suggestive.",
      f"{edu['n_treated']:,} treated and {edu['n_control']:,} comparison households (14 had no education value)."], 13, gap=0.28)
notes(s, "If asked: A propensity score is the chance of having an educated adult given the household's starting traits. Weighting by it makes the two groups similar. The E-value says how strong a hidden cause "
         "would have to be to explain the gap away. Ours is about 1.4, so we call the result suggestive.")

# =========================================================================================================== B6 SHAP by group
s = backup_slide("SHAP importance by factor group")
picture(s, "c09_shap_groups.png", Inches(0.4), Inches(1.5), w=Inches(7.6))
g = k["shap_group_share"]
text(s, Inches(8.3), Inches(1.6), Inches(4.5), Inches(5.3),
     ["**What it shows",
      f"Household structure ({pct(g['Household structure'])}) and assets and housing ({pct(g['Assets & housing'])}) lead. Education is {pct(g['Education'])} and health {pct(g['Health'], 1)}.",
      "**What to keep in mind",
      "A group with many features collects more importance, and features that move together share credit. So a small share does not mean that a group is unimportant."], 14, gap=0.3)
notes(s, "If asked: We added up the importance of all features in each group. Household structure and assets lead. Groups with many features collect more, and features that move together share credit.")

# =========================================================================================================== B7 Glossary
s = backup_slide("Glossary")
table(s, Inches(0.6), Inches(1.5),
      [["Term", "Plain meaning"],
       ["Panel data", "The same households measured more than once"],
       ["Tendulkar line", "The official poverty line: monthly spending per person, set by state and rural or urban area"],
       ["One-hot encoding", "A category becomes one yes/no column per value"],
       ["Out-of-fold prediction", "A prediction made while the household's village was hidden from the model"],
       ["Cut-off (threshold)", "The probability above which we say 'escaped'"],
       ["ROC-AUC", "Chance the model scores a random escaped household above a random stayed-poor one"],
       ["Macro-F1", "The F1-score averaged over both classes"],
       ["SHAP", "A way to split a prediction into the push from each feature"],
       ["95% interval (bootstrap)", "Redraw villages 1,000 times and take the middle 95% of the scores"],
       ["Propensity score", "The chance of a household having a trait (here: an educated adult), given its other traits"],
       ["Data leakage", "Giving the model the answer, or information it would not have in real use"]],
      [3.4, 8.8], 13, row_h=0.45)
notes(s, "Use this glossary to explain any term in one sentence.")

dh.prs.save(dh.OUT)
print("saved", dh.OUT, "slides:", len(prs.slides))
