"""Slide helpers (layout functions) used by build_deck_v3.py."""
import re
from pathlib import Path
ROOT = Path(__file__).resolve().parent  # the project folder
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

FIG = str(ROOT / "outputs" / "figures_ppt") + "/"
OUT = str(ROOT / "IHDS_poverty_escape_presentation_v4.pptx")

BG = RGBColor(0xFA, 0xF7, 0xF0)
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x5B, 0x64, 0x70)
TEAL = RGBColor(0x0F, 0x6E, 0x6A)
ORANGE = RGBColor(0xC2, 0x41, 0x0C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xE3, 0xDE, 0xD3)
TINT = RGBColor(0xE6, 0xF2, 0xF1)
FONT = "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
H = prs.slide_height
BLANK = prs.slide_layouts[6]
count = [0]

SECTIONS = ["Problem and Use Case", "Literature Review", "Dataset and Preprocessing",
            "Methodology and ML Pipeline", "Implementation", "Results and Evaluation",
            "Interpretation of the Model", "Conclusion"]
PRESENTER = {1: "Anhad", 2: "Anhad", 3: "Aman", 4: "Eshaan", 5: "Aman",
             6: "Gokul", 7: "Eshaan", 8: "Anhad"}
DATA_SRC = None  # the data source is named once on the dataset slide and on the references slide


def text(slide, x, y, w, h, content, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, gap=0.45):
    """Text box. Lines starting with '**' are bold headings; URLs become links."""
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    lines = content if isinstance(content, list) else [content]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(size * gap)
        b, c = bold, color
        if line.startswith("**"):
            line, b, c = line[2:], True, TEAL if color == INK else color
        for piece in re.split(r"(https?://\S+)", line):
            if not piece:
                continue
            r = p.add_run()
            r.text = piece
            r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size), b, c, FONT
            if piece.startswith("http"):
                r.hyperlink.address = piece
                r.font.color.rgb = TEAL
    return box


def rect(slide, x, y, w, h, fill=WHITE, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, rnd=0.06):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = rnd
    return s


def new_slide(section, title, source=None):
    count[0] += 1
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    rect(s, 0, 0, Inches(0.16), H, fill=TEAL, line=None, shape=MSO_SHAPE.RECTANGLE)
    if section:
        text(s, Inches(0.6), Inches(0.32), Inches(9), Inches(0.4),
             f"{section}. {SECTIONS[section - 1].upper()}", 12, TEAL, True)
        tag = rect(s, Inches(10.9), Inches(0.32), Inches(2.1), Inches(0.42), fill=TINT, line=None,
                   rnd=0.5)
        tf = tag.text_frame
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        r = tf.paragraphs[0].add_run()
        r.text = "Presenter: " + PRESENTER[section]
        r.font.size, r.font.bold, r.font.name, r.font.color.rgb = Pt(12), True, FONT, TEAL
    text(s, Inches(0.6), Inches(0.65), Inches(12.2), Inches(0.8), title, 26, INK, True)
    text(s, Inches(12.3), Inches(7.0), Inches(0.8), Inches(0.35), str(count[0]), 11, MUTED,
         align=PP_ALIGN.RIGHT)
    if source:
        text(s, Inches(0.6), Inches(7.0), Inches(11.5), Inches(0.35), source, 10, MUTED)
    return s


def picture(slide, name, x, y, w=None, h=None):
    return slide.shapes.add_picture(FIG + name, x, y, width=w, height=h)


def box(slide, x, y, w, h, head, body, size=15, head_color=TEAL, fill=WHITE, gap=0.45):
    rect(slide, x, y, w, h, fill=fill)
    text(slide, x + Inches(0.2), y + Inches(0.12), w - Inches(0.4), Inches(0.45), head, size + 2,
         head_color, True)
    text(slide, x + Inches(0.2), y + Inches(0.58), w - Inches(0.4), h - Inches(0.65), body, size,
         gap=gap)


def note(slide, x, y, w, h, body, size=15):
    """Highlighted observation strip."""
    rect(slide, x, y, w, h, fill=TINT, line=None)
    text(slide, x + Inches(0.2), y, w - Inches(0.4), h, body, size, TEAL, True,
         anchor=MSO_ANCHOR.MIDDLE, gap=0.2)


def table(slide, x, y, rows, col_w, size=13, highlight=None, row_h=0.42):
    t = slide.shapes.add_table(len(rows), len(rows[0]), x, y, Inches(sum(col_w)),
                               Inches(row_h) * len(rows)).table
    for j, cw in enumerate(col_w):
        t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = TEAL if i == 0 else (
                RGBColor(0xFF, 0xF1, 0xE6) if i == highlight else WHITE)
            c.margin_top = c.margin_bottom = Inches(0.05)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = c.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            r = p.add_run()
            r.text = str(val)
            r.font.size, r.font.name = Pt(size), FONT
            r.font.bold = i == 0 or (j == 0)
            r.font.color.rgb = WHITE if i == 0 else INK
    return t


def stat(slide, x, y, w, number, label, color=TEAL):
    rect(slide, x, y, w, Inches(1.35))
    text(slide, x, y + Inches(0.1), w, Inches(0.65), number, 30, color, True, PP_ALIGN.CENTER)
    text(slide, x + Inches(0.1), y + Inches(0.75), w - Inches(0.2), Inches(0.55), label, 13, MUTED,
         align=PP_ALIGN.CENTER)


